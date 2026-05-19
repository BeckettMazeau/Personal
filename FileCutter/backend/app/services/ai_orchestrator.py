import asyncio
import logging
import os
import re
import hashlib
from typing import List, Dict, Any, Optional, Tuple

import aiohttp
import fitz  # PyMuPDF
from pptx import Presentation
from pydantic import BaseModel, Field, ValidationError

from app.core.config import settings
from app.core.utils import get_optimal_batch_size, clean_and_parse_json
from app.core.cache import SQLiteCache

logger = logging.getLogger(__name__)

_TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".log", ".html", ".xml"}
_MAX_TEXT_BYTES = 64 * 1024
_PROTECTED_WORD_RE = re.compile(
    r"(?<![A-Za-z])(homework|report)(?![A-Za-z])", re.IGNORECASE
)
_HOMEWORK_CONTENT_RE = re.compile(r"\bhomework\b", re.IGNORECASE)


class Tier1Verdict(BaseModel):
    id: int
    delete: bool
    confidence: int = Field(ge=1, le=3)
    reasoning: str = ""


class Tier2Verdict(BaseModel):
    delete: bool
    confidence: int = Field(ge=1, le=3)
    reasoning: str = ""


class AIOrchestrator:
    def __init__(
        self,
        lm_studio_url: str = settings.lm_studio_url,
        cache: Optional[SQLiteCache] = None,
    ):
        self.lm_studio_url = lm_studio_url
        self.timeout = aiohttp.ClientTimeout(total=60)
        self.cache = cache if cache is not None else SQLiteCache()

    def _get_file_hash(self, filepath: str) -> str:
        if not os.path.exists(filepath):
            return ""
        hasher = hashlib.sha256()
        try:
            with open(filepath, "rb") as f:
                while chunk := f.read(8192):
                    hasher.update(chunk)
            return hasher.hexdigest()
        except Exception as e:
            logger.error(f"Error hashing {filepath}: {e}")
            return ""

    async def _call_llm(
        self,
        messages: List[Dict[str, str]],
        response_format: Optional[Dict] = None,
    ) -> Optional[str]:
        """Call LM Studio REST API. Returns content string or None."""
        payload = {
            "model": "local-model",
            "messages": messages,
            "temperature": 0.1,
        }
        if response_format:
            payload["response_format"] = response_format

        try:
            async with aiohttp.ClientSession(timeout=self.timeout) as session:
                async with session.post(self.lm_studio_url, json=payload) as response:
                    if response.status >= 400:
                        body = await response.text()
                        logger.error(
                            f"LLM HTTP error {response.status}: {body[:200]!r}"
                        )
                        response.raise_for_status()
                    data = await response.json()
                    return data["choices"][0]["message"]["content"]
        except asyncio.TimeoutError:
            logger.error("LLM request timed out.")
            return None
        except aiohttp.ClientResponseError as e:
            logger.error(f"LLM ClientResponseError status={e.status}: {e.message}")
            return None
        except aiohttp.ClientError as e:
            logger.error(f"Client error during LLM request: {e}")
            return None
        except KeyError as e:
            logger.error(f"Unexpected response format from LLM: {e}")
            return None
        except Exception as e:
            logger.error(f"Unexpected error during LLM request: {e}")
            return None

    async def tier_1_evaluation(
        self, indexed_filenames: List[Tuple[int, str]]
    ) -> List[Dict[str, Any]]:
        """Shallow Pass: analyze filenames by index. Returns validated entries.

        Input contract: list of (id, filename) tuples. The LLM is prompted to
        return JSON objects keyed by `id` (the integer index), not filename,
        so duplicate basenames in different directories survive correctly.
        """
        if not indexed_filenames:
            return []

        valid_ids = {idx for idx, _ in indexed_filenames}

        system_prompt = (
            "Triage files for deletion. For each input line of the form "
            "'<id>: <filename>', return a JSON object in an array with keys: "
            "'id' (int, matching the input id), 'delete' (bool, true if "
            "garbage/download, false if user-made), 'confidence' (int 1-3), "
            "'reasoning' (str). Wrap the array in an object as "
            '{"results": [...]}. ONLY JSON.'
        )
        user_prompt = "Files:\n" + "\n".join(
            f"{idx}: {name}" for idx, name in indexed_filenames
        )

        messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ]

        content = await self._call_llm(
            messages, response_format={"type": "json_object"}
        )
        parsed = clean_and_parse_json(content)

        if isinstance(parsed, dict):
            entries = parsed.get("results")
            if not isinstance(entries, list):
                entries = [parsed] if "id" in parsed else None
        elif isinstance(parsed, list):
            entries = parsed
        else:
            entries = None

        if not entries:
            logger.error(f"Tier-1 LLM returned no usable JSON. Got: {parsed!r}")
            return []

        results: List[Dict[str, Any]] = []
        for entry in entries:
            try:
                verdict = Tier1Verdict.model_validate(entry)
            except ValidationError as e:
                logger.warning(f"Dropping invalid Tier-1 entry {entry!r}: {e}")
                continue
            if verdict.id not in valid_ids:
                logger.warning(
                    f"Dropping Tier-1 entry with unknown id {verdict.id}"
                )
                continue
            results.append(verdict.model_dump())
        return results

    def _extract_text(self, filepath: str, max_words: int = 500) -> str:
        """Extract up to max_words from a supported file type."""
        if not os.path.exists(filepath):
            logger.error(f"File not found: {filepath}")
            return ""

        ext = os.path.splitext(filepath)[1].lower()
        text = ""

        try:
            if ext == ".pdf":
                try:
                    doc = fitz.open(filepath)
                except Exception as e:
                    logger.error(f"Error opening PDF {filepath}: {e}")
                    return ""
                try:
                    for page in doc:
                        text += page.get_text() + " "
                        if len(text.split()) >= max_words:
                            break
                finally:
                    doc.close()
            elif ext in (".pptx", ".ppt"):
                prs = Presentation(filepath)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text += shape.text + " "
                            if len(text.split()) >= max_words:
                                break
                    if len(text.split()) >= max_words:
                        break
            elif ext in _TEXT_EXTENSIONS:
                with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read(_MAX_TEXT_BYTES)
            else:
                return ""
        except Exception as e:
            logger.error(f"Error extracting text from {filepath}: {e}")
            return ""

        words = text.split()
        return " ".join(words[:max_words])

    async def tier_2_evaluation(
        self, filepath: str
    ) -> Tuple[Dict[str, Any], bool]:
        """Deep Pass. Returns (result_dict, cacheable).

        cacheable is False when the LLM call/parse failed; True for genuine
        verdicts and for the "no extractable text" stub (a stable property
        of the file).
        """
        basename = os.path.basename(filepath)
        text = await asyncio.to_thread(self._extract_text, filepath)

        if _PROTECTED_WORD_RE.search(basename) or _HOMEWORK_CONTENT_RE.search(text):
            logger.info(f"Protected item detected: {filepath}")
            return (
                {
                    "filename": basename,
                    "delete": False,
                    "confidence": 1,
                    "reasoning": "Protected item detected by static rule (homework/report).",
                },
                True,
            )

        if not text.strip():
            return (
                {
                    "filename": basename,
                    "delete": False,
                    "confidence": 1,
                    "reasoning": "Could not extract text for analysis.",
                },
                True,
            )

        system_prompt = (
            "Read the extracted text and decide if it is a unique "
            "user-generated report/homework (do not delete) OR generic "
            "downloaded documentation/garbage (can be deleted). Return a "
            "JSON object with keys: 'delete' (bool), 'confidence' (int 1-3), "
            "'reasoning' (str). ONLY JSON."
        )
        user_prompt = f"Extracted Text:\n{text}"

        messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ]

        content = await self._call_llm(
            messages, response_format={"type": "json_object"}
        )
        parsed = clean_and_parse_json(content)

        if isinstance(parsed, dict):
            try:
                verdict = Tier2Verdict.model_validate(parsed)
            except ValidationError as e:
                logger.warning(f"Tier-2 validation failed for {filepath}: {e}")
            else:
                result = verdict.model_dump()
                result["filename"] = basename
                return result, True

        return (
            {
                "filename": basename,
                "delete": False,
                "confidence": 1,
                "reasoning": (
                    "Failed to parse LLM response format."
                    if content
                    else "LLM failed to respond during deep pass."
                ),
            },
            False,
        )

    async def process_files(self, filepaths: List[str]) -> List[Dict[str, Any]]:
        """Run the complete 2-tier evaluation pipeline."""
        final_results: List[Dict[str, Any]] = []
        uncached_filepaths: List[str] = []
        filepath_hash_map: Dict[str, str] = {}

        for filepath in filepaths:
            file_hash = self._get_file_hash(filepath)
            filepath_hash_map[filepath] = file_hash
            cached_result = self.cache.get(file_hash)
            if cached_result:
                final_results.append(cached_result)
            else:
                uncached_filepaths.append(filepath)

        if not uncached_filepaths:
            return final_results

        index_map: Dict[int, str] = {
            idx: fp for idx, fp in enumerate(uncached_filepaths)
        }
        indexed = [
            (idx, os.path.basename(fp)) for idx, fp in index_map.items()
        ]

        batch_size = get_optimal_batch_size()
        batches = [
            indexed[i:i + batch_size] for i in range(0, len(indexed), batch_size)
        ]

        tier_1_results: List[Dict[str, Any]] = []
        for batch in batches:
            batch_results = await self.tier_1_evaluation(batch)
            tier_1_results.extend(batch_results)

        tier_2_tasks = []
        tier_2_items = []

        for item in tier_1_results:
            entry_id = item.get("id")
            if entry_id not in index_map:
                logger.warning(f"Dropping Tier-1 entry with unknown id: {entry_id}")
                continue
            filepath = index_map[entry_id]
            file_hash = filepath_hash_map[filepath]
            filename = os.path.basename(filepath)
            confidence = item.get("confidence", 1)

            output_item = {k: v for k, v in item.items() if k != "id"}
            output_item["filename"] = filename

            if confidence < 3:
                tier_2_tasks.append(self.tier_2_evaluation(filepath))
                tier_2_items.append((filepath, file_hash, output_item))
            else:
                final_results.append(output_item)
                self.cache.set(file_hash, output_item)

        if tier_2_tasks:
            tier_2_outputs = await asyncio.gather(*tier_2_tasks)
            for (filepath, file_hash, original_item), (tier_2_result, cacheable) in zip(
                tier_2_items, tier_2_outputs
            ):
                tier_2_result["reasoning"] = (
                    f"Tier 1 reasoning: {original_item.get('reasoning')}. "
                    f"Tier 2 reasoning: {tier_2_result.get('reasoning')}"
                )
                final_results.append(tier_2_result)
                if cacheable:
                    self.cache.set(file_hash, tier_2_result)

        return final_results
